"""
RubricJudge — Document Loader
Extracts plain text from PDF and DOCX uploads using pdfplumber (primary),
pypdf (fallback), and python-docx.
"""
from __future__ import annotations

import io
import logging
from pathlib import Path

logger = logging.getLogger(__name__)


def extract_text_from_pdf(file_bytes: bytes) -> str:
    """
    Extract text from a PDF file.
    Tries pdfplumber first (better layout handling), then pypdf as fallback.
    """
    text_parts: list[str] = []

    # --- Primary: pdfplumber ---
    try:
        import pdfplumber  # type: ignore

        with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
            for page in pdf.pages:
                page_text = page.extract_text() or ""
                text_parts.append(page_text)
        result = "\n\n".join(text_parts).strip()
        if result:
            return result
    except Exception as exc:
        logger.warning("pdfplumber failed, trying pypdf: %s", exc)

    # --- Fallback: pypdf ---
    try:
        from pypdf import PdfReader  # type: ignore

        reader = PdfReader(io.BytesIO(file_bytes))
        for page in reader.pages:
            text_parts.append(page.extract_text() or "")
        result = "\n\n".join(text_parts).strip()
        if result:
            return result
    except Exception as exc:
        logger.error("pypdf also failed: %s", exc)

    raise ValueError(
        "Could not extract text from the uploaded PDF. "
        "Please try pasting the text directly instead."
    )


def extract_text_from_docx(file_bytes: bytes) -> str:
    """Extract text from a .docx file using python-docx."""
    try:
        from docx import Document  # type: ignore

        doc = Document(io.BytesIO(file_bytes))
        paragraphs: list[str] = []
        for para in doc.paragraphs:
            xml_str = getattr(para._element, "xml", "")
            if "<w:drawing" in xml_str or "<a:graphic" in xml_str:
                paragraphs.append("[Visual Screenshot / Figure Attached]")
            if para.text.strip():
                paragraphs.append(para.text.strip())
        # Also pull text from tables, formatting them as Markdown
        for table in doc.tables:
            for i, row in enumerate(table.rows):
                cleaned_cells = [cell.text.replace('\n', ' ').replace('\r', '').strip() for cell in row.cells]
                row_text = "| " + " | ".join(cleaned_cells) + " |"
                paragraphs.append(row_text)
                if i == 0:
                    separator = "| " + " | ".join(["---"] * len(row.cells)) + " |"
                    paragraphs.append(separator)
        return "\n\n".join(paragraphs).strip()
    except Exception as exc:
        logger.error("python-docx extraction failed: %s", exc)
        raise ValueError(
            "Could not extract text from the uploaded DOCX. "
            "Please try pasting the text directly instead."
        ) from exc


def extract_text_from_upload(filename: str, file_bytes: bytes) -> str:
    """Route extraction by file extension."""
    suffix = Path(filename).suffix.lower()
    if suffix == ".pdf":
        return extract_text_from_pdf(file_bytes)
    elif suffix in {".docx", ".doc"}:
        return extract_text_from_docx(file_bytes)
    elif suffix in {".txt", ".md"}:
        return file_bytes.decode("utf-8", errors="replace")
    else:
        raise ValueError(
            f"Unsupported file type: '{suffix}'. "
            "Please upload PDF, DOCX, or TXT files."
        )
